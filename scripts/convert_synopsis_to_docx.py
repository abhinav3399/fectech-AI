from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "FACTECH_AI_PROJECT_SYNOPSIS.md"
OUTPUT = ROOT / "docs" / "FACTECH_AI_PROJECT_SYNOPSIS.docx"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_run_font(run, name="Tahoma", size=12, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_text(paragraph, text, size=12, bold=False, italic=False):
    for index, part in enumerate(re.split(r"(\*\*.*?\*\*|\*[^*]+\*)", text)):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=size, italic=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, bold=bold, italic=italic)


def configure_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = "Tahoma"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Tahoma")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size in (("Heading 1", 13), ("Heading 2", 13), ("Heading 3", 13)):
        style = document.styles[name]
        style.font.name = "Tahoma"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Tahoma")
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(8)

    if "Synopsis Title" not in [style.name for style in document.styles]:
        style = document.styles.add_style("Synopsis Title", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = document.styles["Synopsis Title"]
    style.font.name = "Tahoma"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Tahoma")
    style.font.size = Pt(16)
    style.font.bold = True
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style.paragraph_format.space_after = Pt(18)


def add_page_number(section):
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend([field_begin, instruction, field_end])
    set_run_font(run, size=10)


def build_document():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_styles(document)
    section = document.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    add_page_number(section)

    title_page = True
    references = False
    for line in lines:
        value = line.strip()
        if not value:
            continue
        if value == "---":
            if title_page:
                document.add_page_break()
                title_page = False
            continue
        if value.startswith("# "):
            paragraph = document.add_paragraph(style="Synopsis Title")
            add_text(paragraph, value[2:].strip(), size=20, bold=True)
            continue
        if value.startswith("## "):
            heading = value[3:].strip()
            references = heading.lower().startswith("8. references")
            paragraph = document.add_paragraph(style="Heading 1")
            add_text(paragraph, heading, size=13, bold=True)
            continue
        if value.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            add_text(paragraph, value[4:].strip(), size=13, bold=True)
            continue
        if value.startswith("**") and value.endswith("**"):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if title_page else WD_ALIGN_PARAGRAPH.LEFT
            add_text(paragraph, value, size=12, bold=True)
            continue
        if re.match(r"^\d+\. ", value):
            paragraph = document.add_paragraph(style="List Number")
            paragraph.paragraph_format.line_spacing = 1.5
            add_text(paragraph, re.sub(r"^\d+\. ", "", value), size=10 if references else 12)
            continue
        if value.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.line_spacing = 1.5
            add_text(paragraph, value[2:], size=10 if references else 12)
            continue
        if value.startswith("|"):
            continue
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if title_page else WD_ALIGN_PARAGRAPH.JUSTIFY
        add_text(paragraph, value, size=10 if references else 12)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()